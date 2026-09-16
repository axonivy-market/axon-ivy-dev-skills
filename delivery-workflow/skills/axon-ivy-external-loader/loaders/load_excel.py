#!/usr/bin/env python3
"""Extract Excel workbook text for AI/requirements analysis.

Requires:
    pip install openpyxl
"""

from __future__ import annotations

import argparse
import glob
import io
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required. Install it with: pip install openpyxl")


XLSX_EXTS = (".xlsx", ".xlsm", ".xltx", ".xltm")


def resolve_inputs(patterns: list[str]) -> list[str]:
    found = []

    for pattern in patterns:
        if os.path.isdir(pattern):
            found += [
                os.path.join(pattern, name)
                for name in os.listdir(pattern)
                if name.lower().endswith(XLSX_EXTS)
            ]
        elif os.path.isfile(pattern):
            found.append(pattern)
        else:
            found += glob.glob(pattern)

    result = []
    seen = set()

    for path in sorted(found):
        if (
            not os.path.isfile(path)
            or not path.lower().endswith(XLSX_EXTS)
            or os.path.basename(path).startswith("~$")
        ):
            continue

        key = os.path.normcase(os.path.abspath(path))
        if key not in seen:
            seen.add(key)
            result.append(path)

    return result


def clean(value, max_chars: int) -> str:
    if value is None:
        return ""

    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\n", " \\n ").replace("\t", " ").strip()

    if len(text) > max_chars:
        return text[:max_chars] + f"... [truncated, {len(text)} chars total]"

    return text


def merge_maps(sheet):
    anchors = {}
    covered = set()

    for rng in sheet.merged_cells.ranges:
        anchor = rng.start_cell.coordinate
        anchors[anchor] = str(rng)

        for row, col in rng.cells:
            coord = sheet.cell(row, col).coordinate
            if coord != anchor:
                covered.add(coord)

    return anchors, covered


def cell_text(coord, cached, formulas, args) -> str:
    cell = cached[coord]
    text = clean(cell.value, args.max_chars)

    # openpyxl does not calculate formulas, so use formula text
    # when no cached result exists.
    if not text and formulas is not None:
        raw = formulas[coord].value
        if isinstance(raw, str) and raw.startswith("="):
            text = clean(raw, args.max_chars)

    metadata = []

    if cell.comment and not args.no_comments:
        comment = clean(cell.comment.text, args.max_chars)
        if comment:
            metadata.append(
                f"comment by {cell.comment.author}: {comment}"
                if cell.comment.author
                else f"comment: {comment}"
            )

    if cell.hyperlink and not args.no_links:
        link = cell.hyperlink.target or cell.hyperlink.location
        if link:
            metadata.append(f"link: {clean(link, args.max_chars)}")

    if metadata:
        meta = "  ".join(f"«{item}»" for item in metadata)
        return f"{text}  {meta}" if text else meta

    return text


def dump_sheet(out, cached, formulas, args):
    anchors, covered = merge_maps(cached)

    rows = cached.max_row or 0
    cols = cached.max_column or 0
    max_row = min(rows, args.max_rows) if args.max_rows else rows
    max_col = min(cols, args.max_cols) if args.max_cols else cols

    out.write(f"=== SHEET: {cached.title} ===\n")
    out.write(f"state: {cached.sheet_state}\n")
    out.write(f"rows: {rows}\n")
    out.write(f"columns: {cols}\n")

    if max_row < rows:
        out.write(f"extracted_rows: {max_row}\n")
    if max_col < cols:
        out.write(f"extracted_columns: {max_col}\n")

    out.write("\n")

    emitted = False

    for row in cached.iter_rows(
        min_row=1,
        max_row=max_row,
        min_col=1,
        max_col=max_col,
    ):
        parts = []

        for cell in row:
            coord = cell.coordinate

            if coord in covered:
                continue

            text = cell_text(coord, cached, formulas, args)
            if not text:
                continue

            if coord in anchors:
                text += f"  «merged {anchors[coord]}»"

            parts.append(
                f"[{coord}] {text}"
                if args.format == "cells"
                else text
            )

        if parts:
            emitted = True
            prefix = "" if args.format == "cells" else f"r{row[0].row}: "
            out.write(prefix + " | ".join(parts) + "\n")

    if not emitted:
        out.write("(no content)\n")

    out.write("\n")


def dump_workbook(out, path, args):
    cached = openpyxl.load_workbook(path, data_only=True)

    try:
        try:
            formulas = openpyxl.load_workbook(path, data_only=False)
        except Exception:
            formulas = None

        out.write("=" * 100 + "\n")
        out.write(f"WORKBOOK: {os.path.abspath(path)}\n")
        out.write(f"SHEETS: {', '.join(cached.sheetnames)}\n")
        out.write("=" * 100 + "\n\n")

        matched = 0

        for sheet in cached.worksheets:
            if args.sheet and not any(
                name.lower() in sheet.title.lower()
                for name in args.sheet
            ):
                continue

            matched += 1
            formula_sheet = (
                formulas[sheet.title]
                if formulas and sheet.title in formulas.sheetnames
                else None
            )

            dump_sheet(out, sheet, formula_sheet, args)

        if args.sheet and not matched:
            out.write(
                f"(no sheet matched {args.sheet}; "
                "see SHEETS above)\n\n"
            )

    finally:
        if formulas is not None:
            formulas.close()
        cached.close()


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract Excel text for requirements analysis."
    )

    parser.add_argument("inputs", nargs="+")
    parser.add_argument("--out")
    parser.add_argument(
        "--format",
        choices=("grid", "cells"),
        default="cells",
    )
    parser.add_argument("--sheet", action="append")
    parser.add_argument("--max-rows", type=int, default=0)
    parser.add_argument("--max-cols", type=int, default=0)
    parser.add_argument("--max-chars", type=int, default=2000)
    parser.add_argument("--no-comments", action="store_true")
    parser.add_argument("--no-links", action="store_true")

    args = parser.parse_args()

    if args.max_rows < 0 or args.max_cols < 0 or args.max_chars <= 0:
        parser.error("row/column limits must be >= 0 and --max-chars > 0")

    paths = resolve_inputs(args.inputs)

    if not paths:
        sys.stderr.write("No Excel files matched.\n")
        return 1

    out = io.StringIO()

    for path in paths:
        try:
            dump_workbook(out, path, args)
        except Exception as exc:
            out.write(
                f"{'=' * 100}\n"
                f"WORKBOOK ERROR: {os.path.abspath(path)}\n"
                f"{type(exc).__name__}: {exc}\n\n"
            )
            sys.stderr.write(f"WARNING: failed to read {path}: {exc}\n")

    text = out.getvalue()

    if args.out:
        path = os.path.abspath(args.out)
        os.makedirs(os.path.dirname(path), exist_ok=True)

        with open(path, "w", encoding="utf-8") as file:
            file.write(text)

        print(
            f"Wrote {len(text)} chars from "
            f"{len(paths)} workbook(s) to {args.out}"
        )
    else:
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")

        sys.stdout.write(text)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())